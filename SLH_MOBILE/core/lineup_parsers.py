# core/lineup_parsers.py
"""
Parsers for different lineup file formats: Excel, CSV, JSON
"""

import json
import csv
from typing import Dict, List, Optional


class LineupParseError(Exception):
    """Raised when parsing fails"""
    pass


class LineupParsers:
    """Collection of parsers for different file formats"""
    
    @staticmethod
    def parse_excel(filepath: str) -> Dict:
        """
        Parse Excel file (.xlsx)
        
        Expected format:
        Sheet "HEMMALAG":
            Row 1: HEMMALAG
            Row 2: Lag | <team name>
            Row 3: Logo | <path or empty>
            Row 5: SPELARE
            Row 6: Nr | Namn | Position
            Row 7+: <data>
            ...
            LEDARE
            Roll | Namn
            <data>
        
        Sheet "BORTALAG": (same format)
        
        Returns:
            Dict with 'home' and 'away' keys
        """
        try:
            import openpyxl
        except ImportError:
            raise LineupParseError("openpyxl not installed. Run: pip install openpyxl --break-system-packages")
        
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
        except Exception as e:
            raise LineupParseError(f"Could not open Excel file: {e}")
        
        result = {}
        
        # Parse HEMMALAG
        if "HEMMALAG" in wb.sheetnames:
            result['home'] = LineupParsers._parse_excel_sheet(wb["HEMMALAG"])
        else:
            raise LineupParseError("Sheet 'HEMMALAG' not found")
        
        # Parse BORTALAG
        if "BORTALAG" in wb.sheetnames:
            result['away'] = LineupParsers._parse_excel_sheet(wb["BORTALAG"])
        else:
            raise LineupParseError("Sheet 'BORTALAG' not found")
        
        return result
    
    @staticmethod
    def _parse_excel_sheet(sheet) -> Dict:
        """Parse one Excel sheet"""
        team_name = ""
        players = []
        staff = []
        
        mode = None  # 'players' or 'staff'
        
        for row_idx, row in enumerate(sheet.iter_rows(min_row=1, values_only=True), start=1):
            if not row or not any(row):
                continue
            
            first_cell = str(row[0]).strip() if row[0] else ""
            
            # Detect sections
            if first_cell == "Lag" and len(row) > 1:
                team_name = str(row[1]).strip() if row[1] else ""
                continue
            
            if first_cell == "SPELARE":
                mode = 'players'
                continue
            
            if first_cell == "LEDARE":
                mode = 'staff'
                continue
            
            # Skip header rows
            if first_cell in ["Nr", "Roll"]:
                continue
            
            # Parse data
            if mode == 'players' and first_cell:
                try:
                    number = str(row[0]).strip() if row[0] else ""
                    name = str(row[1]).strip().upper() if len(row) > 1 and row[1] else ""
                    position = str(row[2]).strip() if len(row) > 2 and row[2] else "F"
                    
                    if number and name:
                        players.append({
                            'number': number,
                            'name': name,
                            'position': position
                        })
                except (ValueError, IndexError):
                    continue
            
            elif mode == 'staff' and first_cell:
                try:
                    role = str(row[0]).strip() if row[0] else ""
                    name = str(row[1]).strip() if len(row) > 1 and row[1] else ""
                    
                    if role and name:
                        staff.append({
                            'role': role,
                            'name': name
                        })
                except (ValueError, IndexError):
                    continue
        
        return {
            'team_name': team_name,
            'players': players,
            'staff': staff
        }
    
    @staticmethod
    def parse_csv(filepath: str) -> Dict:
        """
        Parse CSV file
        
        Expected format:
        HEMMALAG
        Lag,<team name>
        SPELARE
        Nr,Namn,Position
        1,PLAYER NAME,GK
        ...
        LEDARE
        Roll,Namn
        Huvudtränare,NAME
        ...
        BORTALAG
        (same format)
        
        Returns:
            Dict with 'home' and 'away' keys
        """
        try:
            with open(filepath, 'r', encoding='utf-8-sig') as f:
                reader = csv.reader(f)
                rows = list(reader)
        except Exception as e:
            raise LineupParseError(f"Could not read CSV file: {e}")
        
        result = {}
        current_team = None
        current_mode = None
        current_data = {'team_name': '', 'players': [], 'staff': []}
        
        for row in rows:
            if not row or not any(row):
                continue
            
            first_cell = row[0].strip() if row[0] else ""
            
            # Detect team sections
            if first_cell == "HEMMALAG":
                if current_team:  # Save previous team
                    result[current_team] = current_data
                current_team = 'home'
                current_data = {'team_name': '', 'players': [], 'staff': []}
                continue
            
            if first_cell == "BORTALAG":
                if current_team:  # Save previous team
                    result[current_team] = current_data
                current_team = 'away'
                current_data = {'team_name': '', 'players': [], 'staff': []}
                continue
            
            if not current_team:
                continue
            
            # Parse team name
            if first_cell == "Lag" and len(row) > 1:
                current_data['team_name'] = row[1].strip()
                continue
            
            # Detect sections
            if first_cell == "SPELARE":
                current_mode = 'players'
                continue
            
            if first_cell == "LEDARE":
                current_mode = 'staff'
                continue
            
            # Skip headers
            if first_cell in ["Nr", "Roll"]:
                continue
            
            # Parse data
            if current_mode == 'players' and first_cell:
                try:
                    number = row[0].strip()
                    name = row[1].strip().upper() if len(row) > 1 else ""
                    position = row[2].strip() if len(row) > 2 else "F"
                    
                    if number and name:
                        current_data['players'].append({
                            'number': number,
                            'name': name,
                            'position': position
                        })
                except (ValueError, IndexError):
                    continue
            
            elif current_mode == 'staff' and first_cell:
                try:
                    role = row[0].strip()
                    name = row[1].strip() if len(row) > 1 else ""
                    
                    if role and name:
                        current_data['staff'].append({
                            'role': role,
                            'name': name
                        })
                except (ValueError, IndexError):
                    continue
        
        # Save last team
        if current_team:
            result[current_team] = current_data
        
        if 'home' not in result or 'away' not in result:
            raise LineupParseError("Missing HEMMALAG or BORTALAG section")
        
        return result
    
    @staticmethod
    def parse_json(filepath: str) -> Dict:
        """
        Parse JSON file
        
        Expected format:
        {
            "home": {
                "team_name": "Team Name",
                "players": [
                    {"number": "1", "name": "PLAYER NAME", "position": "GK"}
                ],
                "staff": [
                    {"role": "Huvudtränare", "name": "NAME"}
                ]
            },
            "away": {same format}
        }
        
        Returns:
            Dict with 'home' and 'away' keys
        """
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                data = json.load(f)
        except Exception as e:
            raise LineupParseError(f"Could not parse JSON file: {e}")
        
        # Validate structure
        if 'home' not in data or 'away' not in data:
            raise LineupParseError("JSON must contain 'home' and 'away' keys")
        
        for team_key in ['home', 'away']:
            team = data[team_key]
            if 'team_name' not in team:
                team['team_name'] = ""
            if 'players' not in team:
                team['players'] = []
            if 'staff' not in team:
                team['staff'] = []
        
        return data
    
    @staticmethod
    def detect_format(filepath: str) -> str:
        """
        Detect file format from extension
        
        Returns:
            'excel', 'csv', or 'json'
        
        Raises:
            LineupParseError if format not supported
        """
        filepath_lower = filepath.lower()
        
        if filepath_lower.endswith('.xlsx') or filepath_lower.endswith('.xls'):
            return 'excel'
        elif filepath_lower.endswith('.csv'):
            return 'csv'
        elif filepath_lower.endswith('.json'):
            return 'json'
        else:
            raise LineupParseError(f"Unsupported file format. Use .xlsx, .csv, or .json")
    
    @staticmethod
    def parse_file(filepath: str) -> Dict:
        """
        Auto-detect format and parse file
        
        Returns:
            Dict with 'home' and 'away' keys
        """
        format_type = LineupParsers.detect_format(filepath)
        
        if format_type == 'excel':
            return LineupParsers.parse_excel(filepath)
        elif format_type == 'csv':
            return LineupParsers.parse_csv(filepath)
        elif format_type == 'json':
            return LineupParsers.parse_json(filepath)
        else:
            raise LineupParseError(f"Unknown format: {format_type}")
